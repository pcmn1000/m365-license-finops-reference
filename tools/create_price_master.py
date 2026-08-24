from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo


OUTPUT = Path(__file__).parents[1] / "sample-data" / "License-Price-Master.xlsx"

HEADERS = [
    "sku_part_number",
    "license_name",
    "currency",
    "unit_price_monthly",
    "price_type",
    "effective_from",
    "effective_to",
    "billing_term",
    "agreement_id",
    "source_document",
    "owner_upn",
    "approved_by",
    "approved_at_utc",
    "status",
    "reference_only",
    "notes",
]

ROWS = [
    [
        "M365_TEAMS_PREMIUM",
        "Microsoft Teams Premium (add-on)",
        "JPY",
        1499,
        "PublicList",
        "2026-08-03",
        None,
        "AnnualCommitMonthlyPay",
        "DEMO-PUBLIC-LIST",
        "Microsoft public list price JP, annual commitment, tax excluded",
        "admin@m365cpi37573032.onmicrosoft.com",
        "admin@m365cpi37573032.onmicrosoft.com",
        "2026-08-24T00:00:00Z",
        "Approved",
        False,
        "Demo baseline. Replace with the customer contract price.",
    ],
    [
        "MICROSOFT_365_E5_(NO_TEAMS)",
        "Microsoft 365 E5 (no Teams)",
        "JPY",
        7713,
        "PublicList",
        "2026-08-03",
        None,
        "AnnualCommitMonthlyPay",
        "DEMO-PUBLIC-LIST",
        "Microsoft public list price JP, annual commitment, tax excluded",
        "admin@m365cpi37573032.onmicrosoft.com",
        "admin@m365cpi37573032.onmicrosoft.com",
        "2026-08-24T00:00:00Z",
        "Approved",
        False,
        "Demo baseline. Replace with the customer contract price.",
    ],
    [
        "MICROSOFT_365_COPILOT",
        "Microsoft 365 Copilot",
        "JPY",
        4497,
        "PublicList",
        "2026-08-03",
        None,
        "AnnualCommitMonthlyPay",
        "DEMO-PUBLIC-LIST",
        "Microsoft public list price JP, annual commitment, tax excluded",
        "admin@m365cpi37573032.onmicrosoft.com",
        "admin@m365cpi37573032.onmicrosoft.com",
        "2026-08-24T00:00:00Z",
        "Approved",
        False,
        "Demo baseline. Replace with the customer contract price.",
    ],
    [
        "MICROSOFT_TEAMS_ENTERPRISE_NEW",
        "Microsoft Teams Enterprise (base license)",
        "JPY",
        1281,
        "PublicList",
        "2026-08-03",
        None,
        "AnnualCommitMonthlyPay",
        "DEMO-PUBLIC-LIST",
        "Microsoft public list price JP, annual commitment, tax excluded",
        "admin@m365cpi37573032.onmicrosoft.com",
        "admin@m365cpi37573032.onmicrosoft.com",
        "2026-08-24T00:00:00Z",
        "Approved",
        False,
        "Demo baseline. Replace with the customer contract price.",
    ],
]


def add_list_validation(sheet, formula, cells):
    validation = DataValidation(type="list", formula1=formula, allow_blank=False)
    validation.error = "Select a value from the list."
    validation.errorTitle = "Invalid value"
    sheet.add_data_validation(validation)
    validation.add(cells)


def build_workbook():
    workbook = Workbook()
    prices = workbook.active
    prices.title = "LicensePrices"

    header_fill = PatternFill("solid", fgColor="1026B8")
    header_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    input_font = Font(name="Arial", size=10, color="0000FF")
    body_font = Font(name="Arial", size=10, color="000000")
    warning_fill = PatternFill("solid", fgColor="FFF2CC")
    approved_fill = PatternFill("solid", fgColor="E2F0D9")
    thin_gray = Side(style="thin", color="D9D9D9")

    prices.append(HEADERS)
    for row in ROWS:
        prices.append(row)

    for cell in prices[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=thin_gray)

    for row in prices.iter_rows(min_row=2, max_row=prices.max_row):
        for cell in row:
            cell.font = input_font
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(bottom=thin_gray)

    prices.freeze_panes = "A2"
    prices.auto_filter.ref = f"A1:P{prices.max_row}"
    prices.row_dimensions[1].height = 34
    prices.column_dimensions["A"].width = 34
    prices.column_dimensions["B"].width = 42
    prices.column_dimensions["C"].width = 10
    prices.column_dimensions["D"].width = 20
    prices.column_dimensions["E"].width = 14
    prices.column_dimensions["F"].width = 15
    prices.column_dimensions["G"].width = 15
    prices.column_dimensions["H"].width = 28
    prices.column_dimensions["I"].width = 22
    prices.column_dimensions["J"].width = 48
    prices.column_dimensions["K"].width = 38
    prices.column_dimensions["L"].width = 38
    prices.column_dimensions["M"].width = 24
    prices.column_dimensions["N"].width = 14
    prices.column_dimensions["O"].width = 15
    prices.column_dimensions["P"].width = 52

    for cell in prices["D"][1:]:
        cell.number_format = '¥#,##0;[Red](¥#,##0);-'

    table = Table(displayName="LicensePrices", ref=f"A1:P{prices.max_row}")
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    prices.add_table(table)

    add_list_validation(prices, '"JPY,USD,EUR"', f"C2:C1048576")
    add_list_validation(prices, '"Contract,PublicList,Estimate"', f"E2:E1048576")
    add_list_validation(
        prices,
        '"AnnualCommitMonthlyPay,AnnualPrepay,MonthlyCommit"',
        f"H2:H1048576",
    )
    add_list_validation(prices, '"Draft,Approved,Retired"', f"N2:N1048576")
    add_list_validation(prices, '"TRUE,FALSE"', f"O2:O1048576")

    prices.conditional_formatting.add(
        f"A2:P1048576",
        FormulaRule(formula=["$N2=\"Draft\""], fill=warning_fill),
    )
    prices.conditional_formatting.add(
        f"A2:P1048576",
        FormulaRule(formula=["$N2=\"Approved\""], fill=approved_fill),
    )

    instructions = workbook.create_sheet("Instructions")
    instructions.column_dimensions["A"].width = 28
    instructions.column_dimensions["B"].width = 110
    guidance = [
        ("Purpose", "Approved Microsoft 365 license price master consumed by Microsoft Fabric."),
        ("Edit", "Add or update rows inside the LicensePrices Excel table. Do not rename columns or the table."),
        ("Approval", "Only rows with status=Approved are loaded into the production Delta table."),
        ("Effective dates", "Use non-overlapping date ranges for the same SKU and currency."),
        ("Price", "Enter the tax-excluded monthly unit price per assigned user."),
        ("Security", "Do not store client secrets, invoice files, payment data, or confidential contract clauses."),
        ("Demo", "Rows marked PublicList are demonstration values and must be replaced with customer contract prices."),
    ]
    instructions.append(["Item", "Guidance"])
    for item in guidance:
        instructions.append(item)

    for cell in instructions[1]:
        cell.fill = header_fill
        cell.font = header_font
    for row in instructions.iter_rows(min_row=2):
        row[0].font = Font(name="Arial", size=10, bold=True, color="000000")
        row[1].font = body_font
        row[1].alignment = Alignment(wrap_text=True, vertical="top")

    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    return workbook


def validate_workbook(path):
    workbook = load_workbook(path, data_only=False)
    if "LicensePrices" not in workbook.sheetnames:
        raise RuntimeError("LicensePrices sheet is missing.")
    sheet = workbook["LicensePrices"]
    if "LicensePrices" not in sheet.tables:
        raise RuntimeError("LicensePrices Excel table is missing.")
    actual_headers = [cell.value for cell in sheet[1]]
    if actual_headers != HEADERS:
        raise RuntimeError(f"Header mismatch: {actual_headers}")
    errors = []
    for row in sheet.iter_rows(min_row=2, values_only=True):
        record = dict(zip(HEADERS, row))
        if record["status"] == "Approved" and not record["unit_price_monthly"]:
            errors.append(f"Approved row has no price: {record['sku_part_number']}")
    if errors:
        raise RuntimeError("; ".join(errors))


def main():
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    workbook = build_workbook()
    workbook.save(OUTPUT)
    validate_workbook(OUTPUT)
    print(f"Created and validated {OUTPUT}")


if __name__ == "__main__":
    main()